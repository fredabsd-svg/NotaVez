"""Gera as tabelas JSON usadas pelo NotaVez a partir dos anexos OFICIAIS do
Portal Nacional da NFS-e (gov.br/nfse -> Documentação técnica).

Uso (requer openpyxl):
  python3 scripts/atualizar_tabelas.py ANEXO_A.xlsx ANEXO_B.xlsx ANEXO_I.xlsx [ANEXO_VIII.xlsx ANEXO_VII.xlsx]

Anexos VII (cIndOp) e VIII (correlação item LC 116 × NBS × cIndOp × cClassTrib)
ficam na seção RTC da documentação técnica e alimentam o grupo IBS/CBS da DPS.

Saída em src/data/: municipios.json e servicos-nacionais.json, com a versão
dos anexos registrada em src/data/versoes.json. Rode novamente sempre que o
portal publicar novos anexos (regras e leiautes atualizáveis sem mudar código).
"""
import json, os, re, sys
import openpyxl

anexo_a, anexo_b, anexo_i = sys.argv[1:4]
anexo_viii = sys.argv[4] if len(sys.argv) > 4 else None
anexo_vii = sys.argv[5] if len(sys.argv) > 5 else None
out = os.path.join(os.path.dirname(__file__), '..', 'src', 'data')

def rows(path, sheet):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet in wb.sheetnames else next(s for s in wb.worksheets if s.title.strip() == sheet.strip())
    return list(ws.iter_rows(values_only=True))

# Municípios (Anexo A)
# A coluna de UF da planilha oficial vem incompleta; a UF é derivada dos
# 2 primeiros dígitos do código IBGE (código da unidade da federação).
UF_IBGE = {'11': 'RO', '12': 'AC', '13': 'AM', '14': 'RR', '15': 'PA', '16': 'AP', '17': 'TO', '21': 'MA', '22': 'PI',
           '23': 'CE', '24': 'RN', '25': 'PB', '26': 'PE', '27': 'AL', '28': 'SE', '29': 'BA', '31': 'MG', '32': 'ES',
           '33': 'RJ', '35': 'SP', '41': 'PR', '42': 'SC', '43': 'RS', '50': 'MS', '51': 'MT', '52': 'GO', '53': 'DF'}
mun = []
for r in rows(anexo_a, 'TAB.MUN_IBGE')[1:]:
    if r[3]:
        ibge = str(r[3]).zfill(7)
        mun.append({'ibge': ibge, 'nome': str(r[2]).strip(), 'uf': UF_IBGE[ibge[:2]]})
mun.sort(key=lambda m: (m['uf'], m['nome']))

# Lista nacional de serviços (Anexo B) + incidência e grupos exigidos (Anexo I)
itens = {}
servicos = {}
for r in rows(anexo_b, 'LISTA.SERV.NAC.')[1:]:
    cod, item, sub, desd, desc = r[:5]
    if desc is None:
        continue
    desc = re.sub(r'\s+', ' ', str(desc)).strip()
    if cod is None or str(cod) == 'None':
        if str(sub) == '0':
            itens[str(item)] = desc
        continue
    c = str(cod).zfill(6)
    servicos[c] = {'codigo': c, 'descricao': desc, 'item': itens.get(str(item), '')}

for r in rows(anexo_i, 'MUN.INCID_INFO.SERV.')[4:]:
    if not r[0]:
        continue
    c = str(r[0]).zfill(6)
    if c not in servicos:
        continue
    ep, lp, et = (str(x or '') for x in r[2:5])
    servicos[c]['incidencia'] = 'EP' if 'X' in ep else 'LP' if 'X' in lp else 'ET' if 'X' in et else 'OUTRO'
    grupo = str(r[6] or '-').strip()
    servicos[c]['grupoExigido'] = grupo if grupo in ('obra', 'atvEvento') else None

lista = sorted(servicos.values(), key=lambda s: s['codigo'])
json.dump(mun, open(os.path.join(out, 'municipios.json'), 'w'), ensure_ascii=False, separators=(',', ':'))
json.dump(lista, open(os.path.join(out, 'servicos-nacionais.json'), 'w'), ensure_ascii=False, indent=0)
json.dump({'anexoA': os.path.basename(anexo_a), 'anexoB': os.path.basename(anexo_b), 'anexoI': os.path.basename(anexo_i)},
          open(os.path.join(out, 'versoes.json'), 'w'), ensure_ascii=False, indent=2)
print(len(mun), 'municípios;', len(lista), 'serviços')

# ---------- IBS/CBS: correlação do Anexo VIII (células mescladas → blocos) ----------
# Cada item da LC 116 tem blocos de NBS; cada bloco tem as classificações (cClassTrib)
# possíveis e o código indicador da operação (cIndOp). Uma linha com NBS e cClassTrib
# inicia um bloco; linhas só com NBS entram no bloco; linhas só com cClassTrib também.
if anexo_viii:
    so_digitos = lambda v: re.sub(r'\D', '', str(v or ''))
    limpa = lambda v: re.sub(r'\s+', ' ', str(v or '')).strip()
    # Anexo VII: códigos válidos de cIndOp e o que significam.
    cindop = {}
    if anexo_vii:
        for r in rows(anexo_vii, 'cIndOp Public')[1:]:
            c = so_digitos(r[0])
            if c:
                cindop[c.zfill(6)] = {'tipo': limpa(r[1]), 'caracteristica': limpa(r[2]), 'local': limpa(r[3])}
    # A planilha usa células mescladas: expandimos cada mesclagem para que toda linha
    # tenha item, NBS, cIndOp e cClassTrib explícitos, e agrupamos por item e NBS.
    wb8 = openpyxl.load_workbook(anexo_viii, data_only=True)
    ws8 = wb8['tabela geral']
    for faixa in list(ws8.merged_cells.ranges):
        valor = ws8.cell(faixa.min_row, faixa.min_col).value
        ws8.unmerge_cells(str(faixa))
        for linha in range(faixa.min_row, faixa.max_row + 1):
            for coluna in range(faixa.min_col, faixa.max_col + 1):
                ws8.cell(linha, coluna).value = valor
    itens_ibs = {}
    for r in list(ws8.iter_rows(values_only=True))[1:]:
        cod_item, desc_item, nbs, desc_nbs, onerosa, exterior, indop, local, cclass, nome_cclass = (list(r) + [None] * 10)[:10]
        if not cod_item or not nbs:
            continue
        item = so_digitos(cod_item).zfill(4)
        cod_nbs = so_digitos(nbs)
        entrada = next((e for e in itens_ibs.setdefault(item, []) if e['nbs'] == cod_nbs), None)
        if entrada is None:
            entrada = {'nbs': cod_nbs, 'descricao': limpa(desc_nbs), 'cIndOp': [], 'cClassTrib': []}
            itens_ibs[item].append(entrada)
        # Só operações onerosas com adquirente no País (exportação não é suportada nesta versão).
        if indop and str(exterior or 'N').strip().upper().startswith('N') and str(onerosa or 'S').strip().upper().startswith('S'):
            c = so_digitos(indop).zfill(6)
            if c not in [x['codigo'] for x in entrada['cIndOp']]:
                entrada['cIndOp'].append({'codigo': c, 'valido': (c in cindop) if cindop else None})
        if cclass:
            c = so_digitos(cclass).zfill(6)
            if c not in [x['codigo'] for x in entrada['cClassTrib']]:
                entrada['cClassTrib'].append({'codigo': c, 'nome': limpa(nome_cclass)})
    for entradas in itens_ibs.values():
        for e in entradas:
            e['cIndOp'] = [x['codigo'] for x in e['cIndOp'] if x['valido'] is not False]
            e['completo'] = bool(e['cClassTrib']) and bool(e['cIndOp'])
    usados = sorted({c for v in itens_ibs.values() for e in v for c in e['cIndOp']})
    json.dump({'itens': itens_ibs, 'cIndOp': {c: cindop.get(c, {}) for c in usados}},
              open(os.path.join(out, 'ibscbs-correlacao.json'), 'w'), ensure_ascii=False, separators=(',', ':'))
    versoes = json.load(open(os.path.join(out, 'versoes.json')))
    versoes['anexoVIII'] = os.path.basename(anexo_viii)
    if anexo_vii:
        versoes['anexoVII'] = os.path.basename(anexo_vii)
    json.dump(versoes, open(os.path.join(out, 'versoes.json'), 'w'), ensure_ascii=False, indent=2)
    completos = sum(1 for bl in itens_ibs.values() for e in bl if e['completo'])  # noqa
    print(len(itens_ibs), 'itens com correlação IBS/CBS;', completos, 'NBS completas de', sum(len(bl) for bl in itens_ibs.values()))
